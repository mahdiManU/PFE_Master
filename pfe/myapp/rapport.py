from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Pt
import logging

# Set up logging
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

def color_red_in_text(paragraph, words_to_color):
    for run in paragraph.runs:
        text = run.text
        for word in words_to_color:
            if word in text:
                text_parts = text.split(word)
                run.text = text_parts[0]
                for part in text_parts[1:]:
                    colored_run = paragraph.add_run(word)
                    colored_run.font.color.rgb = RGBColor(255, 0, 0)  # Red color
                    paragraph.add_run(part)

def set_table_border(table):
    tbl = table._element
    tblPr = tbl.xpath('./w:tblPr')[0]
    tblBorders = OxmlElement('w:tblBorders')
    for border_name in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        border = OxmlElement(f'w:{border_name}')
        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), '4')
        border.set(qn('w:space'), '0')
        border.set(qn('w:color'), '000000')
        tblBorders.append(border)
    tblPr.append(tblBorders)

def generer_rapport(erreurs, nom_fichier, name):
    name = ".".join(name.split(".")[:-1])
    document = Document()
    erreurS = []
    erreurMI = []
    erreurC = []

    for i in erreurs:
        if "Erreur S" in i:
            erreurS.append(i)
        elif "Erreur MI" in i:
            erreurMI.append(i)
        else:
            erreurC.append(i)

    style = document.styles['Normal']
    style.font.name = 'Calibri'
    
    styleTitre = document.styles.add_style('CenteredTitle', WD_STYLE_TYPE.PARAGRAPH)
    styleTitre.font.size = Pt(24)
    styleTitre.font.bold = True
    styleTitre.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    header_section = document.sections[0]
    header = header_section.header
    header_text = header.paragraphs[0]

    title = document.add_paragraph(f'Rapport de vérification du precesus {name}', style='CenteredTitle')
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    document.add_paragraph("Ce document est un rapport sur la vérification des processus. Il présente les résultats de la vérification concernant trois types d'erreurs : les erreurs dues à des informations manquantes, les erreurs structurelles et les erreurs de contradiction. Chaque type d'erreur est illustré par un tableau contenant l'erreur et son type, la loi associée à cette erreur et la recommandation correspondante.")

    if erreurMI:
        document.add_heading('Erreurs dues à des informations manquantes', level=1)
        table = document.add_table(rows=1, cols=4)
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = 'Erreur'
        hdr_cells[1].text = "Type d'erreur"
        hdr_cells[2].text = "Interpretation legal"
        hdr_cells[3].text = 'Recommandation'

        for cell in hdr_cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.bold = True

        for i in erreurMI:
            val = i.split("#")
            if len(val) >= 4:
                row_cells = table.add_row().cells
                row_cells[0].text = val[0]
                row_cells[1].text = val[1]
                if val[1] == " Condition manquante ":
                    row_cells[2].text = "Cela se produit lorsque la condition n’est mentionnée dans aucun des textes légaux. C'est le problème de l'incomplétude juridique. Précisément l’absence de condition." + val[2]
                if val[1] == " Role manquant ":
                    row_cells[2].text = "Cela se produit lorsque le rôle/l’unité organisationnelle n’est mentionné dans aucun des textes juridiques. C'est le problème de l'incomplétude juridique. Justement l’absence de rôle." + val[2]
                if val[1] == " Vide juridique ":
                     row_cells[2].text = "Cela se produit lors d'un manques d'informations ," + val[2]
                row_cells[3].text = val[3]
                row_cells[0].width = Cm(1)
            else:
                logger.error(f"Unexpected format in error: {i}")

        set_table_border(table)

    if erreurS:
        document.add_heading('Erreurs Structurelles', level=1)
        tableS = document.add_table(rows=1, cols=4)
        hdr_cells = tableS.rows[0].cells
        hdr_cells[0].text = 'Erreur'
        hdr_cells[1].text = "Type d'erreur"
        hdr_cells[2].text = "Interpretation legal"
        hdr_cells[3].text = 'Recommandation'

        for cell in hdr_cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.bold = True

        for i in erreurS:
            val = i.split("#")
            if len(val) >= 4:
                print("fiha 4")
                row_cells = tableS.add_row().cells
                row_cells[0].text = val[0]
                row_cells[1].text = val[1]
                if val[1] == " (Deterministic deadlock)":
                    row_cells[2].text = "En effet, l’utilisation de « soit… soit » dans « établir soit un rapport d’inventaire, soit préparer la lettre de convocation » implique l’exécution de l’un d’entre eux et non des deux. Même ainsi, l’étape suivante nécessitait les résultats des deux « envoyer le rapport et la lettre de convocation ». La déclaration est donc ambiguë. Cela se produit lorsqu’il existe une ambiguïté associée à l’utilisation de « et » et de « ou »." + val[2]
                if val[1] == " (Deterministic lack of synchronization)":
                    row_cells[2].text = "En fait, l'utilisation du « ou » dans la déclaration prête à confusion, dans la mesure où il n'est pas clair s'il s'agit d'un « ou exclusif » ou d'un « ou inclusif ». S'il est considéré comme un « ou inclusif », il sera représenté comme une passerelle inclusive, ce qui ne pose pas de problème. Au contraire, s'il est considéré comme un « ou » exclusif, il sera représenté par une passerelle exclusive. Cela conduira à une situation de manque de synchronisation." + val[2]
                if val[1] == " (Nondeterministic lack of synchronization)":
                    row_cells[2].text = "En fait, l'utilisation du « ou » dans la déclaration prête à confusion, dans la mesure où il n'est pas clair s'il s'agit d'un « ou exclusif » ou d'un « ou inclusif ». Le problème réside dans le fait que le premier 'ou' est considéré comme un 'ou' inclusif, donc il est représenté comme une passerelle inclusive, cependant, le deuxième 'ou' est considéré comme un 'ou' exclusif, donc il est représenté comme passerelle exclusive. Cela conduira à une situation de manque de synchronisation non déterministe." + val[2]
                if val[1] == " (Nondeterministic deadlock)":
                    row_cells[2].text = "La déclaration est ambiguë. Cela se produit lorsqu’il existe une ambiguïté associée à l’utilisation de « et » et de « ou »." + val[2]
                row_cells[3].text = val[3]
                color_red_in_text(row_cells[3].paragraphs[0], ["and // or"])
                row_cells[0].width = Cm(1)
            else:
                logger.error(f"Unexpected format in error: {i}")

        set_table_border(tableS)

    if erreurC:
        document.add_heading('Erreurs dues à des Contradiction', level=1)
        tableC = document.add_table(rows=1, cols=4)
        cont = tableC.rows[0].cells
        cont[0].text = 'Erreur'
        cont[1].text = "Type d'erreur"
        cont[2].text = "Interpretation legal"
        cont[3].text = 'Recommandation'

        for cell in cont:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.bold = True

        for i in erreurC:
            val = i.split("#")
            if len(val) >= 4:
                row_cells = tableC.add_row().cells
                row_cells[0].text = val[0]
                row_cells[1].text = val[1]
                row_cells[2].text = "Contradiction between legal texts occurs when there is no consistency between them. Thus, a statement in one text totally or partially contradicts a statement in another text. Two forms are considered: the negation and antonymy." + val[2]
                row_cells[3].text = val[3]
                row_cells[0].width = Cm(1)
            else:
                logger.error(f"Unexpected format in error: {i}")

        set_table_border(tableC)

    document.add_page_break()
    document.save(nom_fichier)
    logger.info(f"Report saved at: {nom_fichier}")

# rapportEng.py
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Pt
import logging

# Set up logging
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

def color_red_in_text(paragraph, words_to_color):
    for run in paragraph.runs:
        text = run.text
        for word in words_to_color:
            if word in text:
                text_parts = text.split(word)
                run.text = text_parts[0]
                for part in text_parts[1:]:
                    colored_run = paragraph.add_run(word)
                    colored_run.font.color.rgb = RGBColor(255, 0, 0)  # Red color
                    paragraph.add_run(part)

def set_table_border(table):
    tbl = table._element
    tblPr = tbl.xpath('./w:tblPr')[0]
    tblBorders = OxmlElement('w:tblBorders')
    for border_name in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        border = OxmlElement(f'w:{border_name}')
        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), '4')
        border.set(qn('w:space'), '0')
        border.set(qn('w:color'), '000000')
        tblBorders.append(border)
    tblPr.append(tblBorders)

def generer_rapportENG(erreurs, nom_fichier, name):
    name = ".".join(name.split(".")[:-1])
    document = Document()
    erreurS = []
    erreurMI = []
    erreurC = []

    for i in erreurs:
        if "Error S" in i:
            erreurS.append(i)
        elif "Error MI" in i:
            erreurMI.append(i)
        else:
            erreurC.append(i)

    style = document.styles['Normal']
    style.font.name = 'Calibri'
    
    styleTitre = document.styles.add_style('CenteredTitle', WD_STYLE_TYPE.PARAGRAPH)
    styleTitre.font.size = Pt(24)
    styleTitre.font.bold = True
    styleTitre.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    header_section = document.sections[0]
    header = header_section.header
    header_text = header.paragraphs[0]

    title = document.add_paragraph(f'Precesus Verification Report {name}', style='CenteredTitle')
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    document.add_paragraph("This document is a process audit report. It presents the verification results for three types of errors: errors due to missing information, structural errors and contradiction errors. Each type of error is illustrated by a table containing the error and its type, the law associated with this error and the corresponding recommendation.")

    if erreurMI:
        document.add_heading('Errors due to missing information', level=1)
        table = document.add_table(rows=1, cols=4)
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = 'Error'
        hdr_cells[1].text = "Error type"
        hdr_cells[2].text = "Legal interpretation"
        hdr_cells[3].text = 'Recommendation'

        for cell in hdr_cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.bold = True

        for i in erreurMI:
            val = i.split("#")
            if len(val) >= 4:
                row_cells = table.add_row().cells
                row_cells[0].text = val[0]
                row_cells[1].text = val[1]
                if val[1] == " Missing condition ":
                    row_cells[2].text = "This happens when the condition is not mentioned in any of the legal texts. This is the problem of legal incompleteness. Precisely the absence of conditions." + val[2]
                if val[1] == " Missing role ":
                    row_cells[2].text = "This happens when the role/organizational unit is not mentioned in any of the legal texts. This is the problem of legal incompleteness. Precisely the absence of a role." + val[2]
                if val[1] == " Legal vacuum ":
                     row_cells[2].text = "This happens when there is a lack of information," + val[2]
                row_cells[3].text = val[3]
                row_cells[0].width = Cm(1)
            else:
                logger.error(f"Unexpected format in error: {i}")

        set_table_border(table)

    if erreurS:
        document.add_heading('Structural Errors', level=1)
        tableS = document.add_table(rows=1, cols=4)
        hdr_cells = tableS.rows[0].cells
        hdr_cells[0].text = 'Error'
        hdr_cells[1].text = "Error type"
        hdr_cells[2].text = "Legal interpretation"
        hdr_cells[3].text = 'Recommendation'

        for cell in hdr_cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.bold = True

        for i in erreurS:
            val = i.split("#")
            if len(val) >= 4:
                print("fiha 4")
                row_cells = tableS.add_row().cells
                row_cells[0].text = val[0]
                row_cells[1].text = val[1]
                if val[1] == " (Deterministic deadlock)":
                    row_cells[2].text = "Indeed, the use of “either…or” in “establish either an inventory report or prepare the letter of invitation” implies the execution of one of them and not both. Even so, the next step required the results of both “send report and summons letter”. The statement is therefore ambiguous. This occurs when there is ambiguity associated with the use of 'and' and 'or'." + val[2]
                if val[1] == " (Deterministic lack of synchronization)":
                    row_cells[2].text = "In fact, the use of 'or' in the statement is confusing, as it is not clear whether it is an 'exclusive or' or an 'inclusive or'. If it is considered an 'inclusive or', it will be represented as an inclusive gateway, which is not a problem. On the contrary, if it is considered an exclusive 'or', it will be represented by an exclusive gateway. This will lead to a situation of lack of synchronization." + val[2]
                if val[1] == " (Nondeterministic lack of synchronization)":
                    row_cells[2].text = "In fact, the use of 'or' in the statement is confusing, as it is not clear whether it is an 'exclusive or' or an 'inclusive or'. The problem lies in the fact that the first 'or' is considered an inclusive 'or', so it is represented as an inclusive gateway, however, the second 'or' is considered an exclusive 'or', therefore it is represented as an exclusive gateway. This will lead to a non-deterministic lack of synchronization situation." + val[2]
                if val[1] == " (Nondeterministic deadlock)":
                    row_cells[2].text = "The statement is ambiguous. This occurs when there is ambiguity associated with the use of 'and' and 'or'." + val[2]
                row_cells[3].text = val[3]
                color_red_in_text(row_cells[3].paragraphs[0], ["and // or"])
                row_cells[0].width = Cm(1)
            else:
                logger.error(f"Unexpected format in error: {i}")

        set_table_border(tableS)

    if erreurC:
        document.add_heading('Erreurs dues à des Contradiction', level=1)
        tableC = document.add_table(rows=1, cols=4)
        cont = tableC.rows[0].cells
        cont[0].text = 'Error'
        cont[1].text = "Error type"
        cont[2].text = "Legal interpretation"
        cont[3].text = 'Recommendation'

        for cell in cont:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.bold = True

        for i in erreurC:
            val = i.split("#")
            if len(val) >= 4:
                row_cells = tableC.add_row().cells
                row_cells[0].text = val[0]
                row_cells[1].text = val[1]
                row_cells[2].text = "Contradiction between legal texts occurs when there is no consistency between them. Thus, a statement in one text totally or partially contradicts a statement in another text. Two forms are considered: the negation and antonymy." + val[2]
                row_cells[3].text = val[3]
                row_cells[0].width = Cm(1)
            else:
                logger.error(f"Unexpected format in error: {i}")

        set_table_border(tableC)

    document.add_page_break()
    document.save(nom_fichier)
    logger.info(f"Report saved at: {nom_fichier}")
